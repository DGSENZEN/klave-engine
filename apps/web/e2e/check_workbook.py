"""Assert independently specified quantities, prices, and evidence in a download."""

import json
import math
import sys
from pathlib import Path

from openpyxl import load_workbook


def check_workbook(path: Path, expected_path: Path) -> None:
    expected = json.loads(expected_path.read_text())
    workbook = load_workbook(path, data_only=True)
    try:
        assert {"Presupuesto", "Generadores", "Lo que el plano no dio"} <= set(workbook.sheetnames)
        rows = list(workbook["Presupuesto"].values)
        headers = list(rows[0])
        code = expected.get("export_code", expected["concept_code"])
        selected = [row for row in rows[1:] if row[headers.index("Clave")] == code]
        assert len(selected) == 1, "Expected exactly one explicitly priced office concept"
        row = selected[0]
        assert row[headers.index("Unidad")] == expected["unit"]
        for header, value in (
            ("Cantidad", expected["quantity"]),
            ("P.U. (CD)", expected["unit_price"]),
            ("Importe", expected["quantity"] * expected["unit_price"]),
        ):
            actual = row[headers.index(header)]
            assert isinstance(actual, (int, float)), "Expected a priced numeric cell: " + header
            assert math.isclose(actual, value, abs_tol=0.000001), "Unexpected " + header
        assert workbook["Generadores"].max_row >= 2
        assert workbook["Lo que el plano no dio"].max_row >= 1
    finally:
        workbook.close()


if __name__ == "__main__":
    check_workbook(Path(sys.argv[1]), Path(sys.argv[2]))
    print("Downloaded workbook quantities, prices, and disclosures verified.")
