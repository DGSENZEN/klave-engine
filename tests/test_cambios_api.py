"""Las revisiones de un proyecto y qué cambió entre dos, por API, con su
libro de aditivas y deductivas y la honestidad de «mismo plano»."""

import io
import json

from klave_engine.costing.catalog_store import get_catalog_store
from klave_engine.costing.models import CostingOverrides
from klave_engine.costing.recompute import recompute_and_persist
from klave_engine.detection.results import DetectionType, make_detection
from openpyxl import load_workbook


def _cols(xs):
    out = []
    for i, x in enumerate(xs):
        d = make_detection(f"c{i}", DetectionType.column_tag, "EST", (x, 0, x + 0.3, 0.4), 0.9,
                           [f"c{i}"], "tag", [], {"section_cm": "30x40", "section_source": "cuadro",
                                                   "section_area_du2": 0.0}, "e.dxf")
        out.append(d.model_copy(update={"mark": "C-1"}).model_dump(mode="json"))
    return out


def _run(processed, run_id, dets, engine, inputs, store, pid, root):
    folder = processed / "runs" / run_id
    folder.mkdir(parents=True)
    (folder / "detections.json").write_text(json.dumps(dets))
    (folder / "drawing_units.json").write_text(json.dumps(
        {"unit": "m", "source": "declared", "confidence": 1.0}))
    (folder / "engine.json").write_text(json.dumps(
        {"version": "1", "fingerprint": engine, "processed_at": f"2026-10-0{run_id[-1]}T00:00:00"}))
    if inputs is not None:
        (folder / "inputs.json").write_text(json.dumps({"files": inputs}))
    recompute_and_persist(folder, folder, root / "reports", pid, CostingOverrides(),
                          catalog_store=store)
    # El reporte persistido de una corrida vive en su carpeta.
    override = folder / "cost_report_override.json"
    if override.exists():
        override.rename(folder / "cost_report.json")


def _client(data_dir, monkeypatch, *, same_drawing):
    from fastapi.testclient import TestClient
    from klave_engine.common import config as config_module

    from apps.api.main import create_app

    pid = "obra-cambios"
    root = data_dir / "projects" / pid
    processed = root / "processed"
    processed.mkdir(parents=True)
    (processed / "project_manifest.json").write_text(json.dumps({
        "project_id": pid, "project_name": "Obra", "root_path": str(root), "drawings": [],
    }))
    (data_dir / "projects_registry.json").write_text(json.dumps({pid: str(root)}))
    store = get_catalog_store(data_dir)
    _run(processed, "run_1", _cols([0, 5, 10]), "eng", {"E.dwg": "aaa"}, store, pid, root)
    _run(processed, "run_2", _cols([0, 6, 10, 20]), "eng",
         {"E.dwg": "aaa" if same_drawing else "bbb"}, store, pid, root)
    (processed / "active_run.json").write_text(json.dumps(
        {"run_id": "run_2", "artifact_dir": "runs/run_2", "reports_dir": "runs/run_2/reports"}))
    config_module.get_settings.cache_clear()
    return TestClient(create_app()), pid


def test_revisions_changes_labels_and_xlsx(data_dir, monkeypatch):
    client, pid = _client(data_dir, monkeypatch, same_drawing=False)
    revs = client.get(f"/projects/{pid}/revisiones").json()["revisiones"]
    assert [r["run_id"] for r in revs] == ["run_1", "run_2"]
    assert revs[1]["active"] and revs[1]["mismo_plano_que_anterior"] is False

    cambios = client.get(f"/projects/{pid}/cambios")
    assert cambios.status_code == 200, cambios.text
    body = cambios.json()
    assert body["antes"] == "run_1" and body["despues"] == "run_2"
    assert body["resumen"] == {"movido": 1, "agregado": 1}
    assert body["mismo_plano"] is False
    est = next(c for c in body["conceptos"] if c["concept_code"] == "EST-001")
    assert est["diferencia"] > 0 and est["elementos"]["agregado"] == 1

    named = client.put(f"/projects/{pid}/revisiones/run_2", json={"label": "Rev C"})
    assert named.status_code == 200
    assert client.get(f"/projects/{pid}/cambios").json()["despues_label"] == "Rev C"
    assert client.put(f"/projects/{pid}/revisiones/../../x", json={"label": "x"}).status_code in (
        404, 405)
    assert client.get(f"/projects/{pid}/cambios?antes=run_1&despues=run_9").status_code == 404
    assert client.get(f"/projects/{pid}/cambios?antes=run_1&despues=run_1").status_code == 422

    xlsx = client.get(f"/projects/{pid}/cambios.xlsx")
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["Aditivas y deductivas", "Elementos"]
    claves = [c.value for c in book["Aditivas y deductivas"]["A"][4:] if c.value]
    assert not any(str(c).startswith("EST-") for c in claves)


def test_same_drawing_is_said(data_dir, monkeypatch):
    client, pid = _client(data_dir, monkeypatch, same_drawing=True)
    assert client.get(f"/projects/{pid}/cambios").json()["mismo_plano"] is True
