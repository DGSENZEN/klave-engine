"""La liga para compartir: dueño la crea y la revoca; quien la abre ve plano y
generadores sin cuenta; caducada, revocada o inventada responde lo mismo."""

import io
import json

from openpyxl import load_workbook

from apps.api.auth.middleware import OPEN_PREFIXES, _required_project_role
from tests.test_variantes_api import _project


def _client(data_dir, monkeypatch):
    from fastapi.testclient import TestClient
    from klave_engine.common import config as config_module

    from apps.api.main import create_app

    monkeypatch.setenv("KLAVE_USERS_DATABASE_URL", "postgresql://nobody@127.0.0.1:1/none")
    config_module.get_settings.cache_clear()
    return TestClient(create_app())


def test_policy_is_owner_to_create_and_token_to_read():
    assert _required_project_role(["projects", "p", "compartir"], "POST") == "owner"
    assert _required_project_role(["projects", "p", "compartir"], "GET") == "owner"
    assert "/compartido/" in OPEN_PREFIXES
    # El prefijo abierto termina en «/»: «/compartidos» no se cuela.
    assert not any("/compartidos".startswith(p) for p in OPEN_PREFIXES if p != "/compartido/")


def test_share_link_lifecycle(data_dir, monkeypatch):
    pid, processed = _project(data_dir)
    (processed / "normalized_entities.json").write_text("[]")
    client = _client(data_dir, monkeypatch)

    created = client.post(f"/projects/{pid}/compartir", json={"days": 3},
                          headers={"X-Actor": "Diego"})
    assert created.status_code == 201, created.text
    token = created.json()["token"]
    assert len(token) >= 32 and created.json()["url"].endswith(f"/compartido/{token}")
    assert client.get(f"/projects/{pid}/compartir").json()["links"][0]["active"] is True

    meta = client.get(f"/compartido/{token}")
    assert meta.status_code == 200 and meta.json()["project_name"] == "Obra"
    geometry = client.get(f"/compartido/{token}/geometry")
    assert geometry.status_code == 200 and "detections" in geometry.json()
    xlsx = client.get(f"/compartido/{token}/generadores.xlsx")
    assert xlsx.status_code == 200
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["Generadores", "Lo que el plano no dio"]
    # Ningún peso en lo compartido.
    cells = [c.value for ws in book.worksheets for row in ws.iter_rows() for c in row]
    assert not any(isinstance(v, str) and "$" in v for v in cells)

    assert client.get("/compartido/no-es-un-token-de-verdad-123").status_code == 404
    assert client.delete(f"/projects/{pid}/compartir/{token}").status_code == 200
    gone = client.get(f"/compartido/{token}")
    assert gone.status_code == 404 and "caducó" in gone.json()["detail"]["message"]

    again = client.post(f"/projects/{pid}/compartir", json={}).json()["token"]
    links = json.loads((data_dir / "share_links.json").read_text())
    links[again]["expires_at"] = "2020-01-01T00:00:00+00:00"
    (data_dir / "share_links.json").write_text(json.dumps(links))
    assert client.get(f"/compartido/{again}/geometry").status_code == 404
    assert client.post(f"/projects/{pid}/compartir", json={"days": 400}).status_code == 422
