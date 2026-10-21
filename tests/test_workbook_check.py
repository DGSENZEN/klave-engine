import importlib.util
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

SCRIPT = Path(__file__).resolve().parents[1] / "apps/web/e2e/check_workbook.py"


def checker():
    if not SCRIPT.exists():
        pytest.fail("downloaded workbook acceptance checker is missing")
    spec = importlib.util.spec_from_file_location("workbook_check", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("price,passes", [(1000, True), (0, False), (999, False)])
def test_workbook_checks_independent_values(tmp_path, price, passes):
    module = checker()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Presupuesto"
    sheet.append(["Clave", "Concepto", "Unidad", "Cantidad", "P.U. (CD)", "Importe"])
    sheet.append(["PILOT-SLAB", "Synthetic slab area", "M2", 12, price, 12 * price])
    workbook.create_sheet("Generadores").append(["Hoja", "Evidencia"])
    workbook["Generadores"].append(["S-101", "4 × 3"])
    workbook.create_sheet("Lo que el plano no dio").append(["Omisiones declaradas"])
    path = tmp_path / "budget.xlsx"
    workbook.save(path)
    expected = tmp_path / "expected.json"
    expected.write_text(json.dumps({
        "concept_code": "EST-016", "export_code": "PILOT-SLAB",
        "unit": "M2", "quantity": 12, "unit_price": 1000,
    }))
    if passes:
        module.check_workbook(path, expected)
    else:
        with pytest.raises(AssertionError):
            module.check_workbook(path, expected)
